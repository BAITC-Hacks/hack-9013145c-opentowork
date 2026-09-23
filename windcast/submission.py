"""Черновики заявки на продажу и корректировок для РФЦ по прогнозу агента.

Основание — Правила организации и функционирования оптового рынка электрической энергии
(приказ Минэнерго РК, V1500010531):
  • приложение 7 — форма «Заявка на продажу»: 24 часовых значения в МВт, до тысячных;
  • п. 51 — ВИЭ с долгосрочным договором подаёт одну заявку единому закупщику
    до 08:00 по времени Астаны суток, предшествующих операционным;
  • п. 97–99 — корректировка не позднее чем за 2 часа до наступления часа.

Документ — ЧЕРНОВИК: настоящая заявка подаётся в системе балансирующего рынка с ЭЦП.
Числа считает код из прогноза агента; LLM к объёмам в заявке не допускается.

Допущения, которые печатаются в самом документе:
  • номинал турбины RATED_MW_PER_TURBINE — в данных кейса не указан;
  • собственное потребление станции = 0;
  • «01:00» в форме = интервал 00:00–01:00 (метка конца часа); у нас метка — начало часа.
"""

from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass, field
from pathlib import Path

import pandas as pd

from windcast.config import ARTIFACTS, FORECASTS_DIR, SCADA_UTC_OFFSET_H, TURBINES

BIDS_DIR = ARTIFACTS / "bids"
FONTS = Path(__file__).parent / "assets" / "fonts"

RATED_MW_PER_TURBINE = 2.5  # допущение: номинал в данных не указан
ASTANA_UTC_OFFSET_H = SCADA_UTC_OFFSET_H  # время Астаны = UTC+5, как и время SCADA
SUBMISSION_DEADLINE_LOCAL_H = 8  # п. 51: до 08:00 суток D−1
CORRECTION_LEAD_H = 2  # п. 97, 99: не позднее чем за 2 часа до часа
CORRECTION_MIN_SHARE = 0.10  # корректируем, если сдвиг ≥ 10% установленной мощности

DEFAULT_SENDER = "<наименование энергопроизводящей организации> — ВЭС, турбины T1, T2"
COUNTERPARTY = (
    "ТОО «Расчетно-финансовый центр по поддержке возобновляемых источников энергии» "
    "(единый закупщик)"
)
DRAFT_BANNER = (
    "ЧЕРНОВИК — не является поданной заявкой. Подача осуществляется в системе "
    "балансирующего рынка с ЭЦП организации."
)


@dataclass
class HourBid:
    hour: int  # 1..24, метка конца часа как в форме
    interval_local: str  # «00:00–01:00»
    utc_start: str
    mw: float  # в заявку (P50)
    p10_mw: float
    p90_mw: float
    wind_ms: float
    wind_spread_ms: float
    horizon_h: int
    icing: bool


@dataclass
class Correction:
    decided_at_local: str
    weather_run_origin: str
    hour: int
    interval_local: str
    direction: str  # «вверх» / «вниз»
    volume_mw: float
    was_mw: float
    new_mw: float
    submit_before_local: str
    partial_allowed: bool = True


@dataclass
class Bid:
    operational_day: str
    deadline_local: str
    prepared_at_local: str
    sender: str
    counterparty: str
    operation: str
    installed_mw: float
    rated_mw_per_turbine: float
    forecast_id: str
    model_version: str
    weather_provider: str
    weather_run: str
    hours: list[HourBid]
    total_mwh: float
    corrections: list[Correction] = field(default_factory=list)
    risks: list[str] = field(default_factory=list)
    assumptions: list[str] = field(default_factory=list)
    agent_summary: str = ""


def _utc(ts: str) -> pd.Timestamp:
    return pd.Timestamp(ts.rstrip("Z"))


def _local(ts: pd.Timestamp) -> pd.Timestamp:
    return ts + pd.Timedelta(hours=ASTANA_UTC_OFFSET_H)


def _load_day(day: pd.Timestamp) -> dict | None:
    path = FORECASTS_DIR / f"{day.date()}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def _hour_slots(day: pd.Timestamp) -> list[tuple[int, pd.Timestamp, str]]:
    """(номер часа 1..24, начало часа в UTC, «ЧЧ:00–ЧЧ:00» по Астане) для суток D."""
    slots = []
    for k in range(1, 25):
        local_start = day + pd.Timedelta(hours=k - 1)
        utc_start = local_start - pd.Timedelta(hours=ASTANA_UTC_OFFSET_H)
        label = f"{k - 1:02d}:00–{k:02d}:00"
        slots.append((k, utc_start, label))
    return slots


def _fmt(x: float) -> str:
    """Формат формы: три знака, запятая."""
    return f"{x:.3f}".replace(".", ",")


def build_bid(
    operational_day: str | pd.Timestamp,
    sender: str = DEFAULT_SENDER,
    rated_mw_per_turbine: float = RATED_MW_PER_TURBINE,
) -> Bid:
    day = pd.Timestamp(operational_day).normalize()
    issue_day = day - pd.Timedelta(days=1)
    doc = _load_day(issue_day)
    if doc is None:
        raise FileNotFoundError(
            f"нет прогона агента за {issue_day.date()} — python -m windcast test-period"
        )
    installed = rated_mw_per_turbine * len(TURBINES)
    runs = [v for v in doc["versions"] if "skipped" not in v]
    deadline_utc = issue_day + pd.Timedelta(hours=SUBMISSION_DEADLINE_LOCAL_H - ASTANA_UTC_OFFSET_H)
    # В заявку идёт последняя опубликованная версия, выпущенная до срока подачи.
    eligible = [v for v in runs if v["published"] and _utc(v["forecast_origin"]) <= deadline_utc]
    if not eligible:
        raise ValueError(f"нет прогноза, выпущенного до срока подачи {deadline_utc} UTC")
    base = eligible[-1]
    by_time = {_utc(p["forecast_for"]): p for p in base["predictions"]}

    hours = []
    for k, utc_start, label in _hour_slots(day):
        p = by_time.get(utc_start)
        if p is None:
            raise ValueError(f"прогноз {base['forecast_id']} не покрывает час {label} {day.date()}")
        hours.append(
            HourBid(
                hour=k,
                interval_local=label,
                utc_start=str(utc_start),
                mw=round(p["p50"] * installed, 3),
                p10_mw=round(p["p10"] * installed, 3),
                p90_mw=round(p["p90"] * installed, 3),
                wind_ms=p["wind_speed"],
                wind_spread_ms=p.get("wind_spread") or 0.0,
                horizon_h=p["horizon_h"],
                icing=bool(p.get("icing_risk")),
            )
        )

    # Риски — только по 24 часам операционных суток: сводка агента описывает все 48 ч.
    facts = base.get("facts", {})
    risks = []
    disagree = [h for h in hours if h.wind_spread_ms > 2.5]
    if disagree:
        risks.append(f"погодные модели расходятся больше 2.5 м/с в {len(disagree)} ч")
    icing = [h for h in hours if h.icing]
    if icing:
        risks.append(
            f"риск обледенения в {len(icing)} ч (t −8…+1 °C, влажность ≥ 90%): "
            f"{icing[0].interval_local} … {icing[-1].interval_local}"
        )
    day_hours = {h.utc_start for h in hours}
    for r in facts.get("analysis", {}).get("ramps", []):
        if str(_utc(r["time"])) in day_hours:
            word = "рост" if r["direction"] == "up" else "спад"
            risks.append(
                f"резкий {word} на {abs(r['delta']) * installed:.2f} МВт за 3 ч к "
                f"{_local(_utc(r['time'])):%H:%M}"
            )
    wide = [h for h in hours if (h.p90_mw - h.p10_mw) > 0.45 * installed]
    if wide:
        risks.append(
            f"низкая уверенность в {len(wide)} ч (коридор P10–P90 шире 45% установленной мощности)"
        )

    bid = Bid(
        operational_day=str(day.date()),
        deadline_local=f"{issue_day:%d.%m.%Y} {SUBMISSION_DEADLINE_LOCAL_H:02d}:00",
        prepared_at_local=f"{_local(_utc(base['forecast_origin'])):%d.%m.%Y %H:%M}",
        sender=sender,
        counterparty=COUNTERPARTY,
        operation="Продажа",
        installed_mw=installed,
        rated_mw_per_turbine=rated_mw_per_turbine,
        forecast_id=base["forecast_id"],
        model_version=base["model_version"],
        weather_provider=base["weather_provider"],
        weather_run=base["weather_run"],
        hours=hours,
        total_mwh=round(sum(h.mw for h in hours), 3),
        risks=risks,
        assumptions=[
            f"Номинал турбины {rated_mw_per_turbine:g} МВт — допущение: в данных кейса не указан; "
            f"установленная мощность {installed:g} МВт.",
            "Собственное потребление станции принято равным 0.",
            "Час «01:00» формы = интервал 00:00–01:00 по времени Астаны (UTC+5). "
            "Нумерацию часов уточнить у РФЦ.",
            "В заявку внесена медиана прогноза (P50); коридор P10–P90 — справочно.",
        ],
    )
    bid.corrections = build_corrections(bid, base, issue_day)
    return bid


def build_corrections(bid: Bid, base: dict, issue_day: pd.Timestamp) -> list[Correction]:
    """Черновики корректировок из ревизий агента (п. 97–99).

    Ревизии берутся из прогонов суток D−1 и D, выпущенных после базовой версии.
    Час доступен для корректировки, если до его начала не меньше 2 часов.
    """
    installed = bid.installed_mw
    threshold = CORRECTION_MIN_SHARE * installed
    schedule = {pd.Timestamp(h.utc_start): h.mw for h in bid.hours}
    slot = {pd.Timestamp(h.utc_start): h for h in bid.hours}
    base_origin = _utc(base["forecast_origin"])
    revisions = []
    for d in (issue_day, issue_day + pd.Timedelta(days=1)):
        doc = _load_day(d)
        for v in doc["versions"] if doc else []:
            if "skipped" in v or not v["published"]:
                continue
            if _utc(v["forecast_origin"]) > base_origin:
                revisions.append(v)
    revisions.sort(key=lambda v: v["forecast_origin"])

    lead = pd.Timedelta(hours=CORRECTION_LEAD_H)
    out = []
    for v in revisions:
        decided = _utc(v["forecast_origin"])
        for p in v["predictions"]:
            t = _utc(p["forecast_for"])
            if t not in schedule or t < decided + pd.Timedelta(hours=CORRECTION_LEAD_H):
                continue
            new = round(p["p50"] * installed, 3)
            delta = new - schedule[t]
            if abs(delta) < threshold:
                continue
            h = slot[t]
            out.append(
                Correction(
                    decided_at_local=f"{_local(decided):%d.%m %H:%M}",
                    weather_run_origin=v["forecast_origin"],
                    hour=h.hour,
                    interval_local=h.interval_local,
                    direction="вверх" if delta > 0 else "вниз",
                    volume_mw=round(abs(delta), 3),
                    was_mw=schedule[t],
                    new_mw=new,
                    submit_before_local=f"{_local(t - lead):%d.%m %H:%M}",
                )
            )
            schedule[t] = new
    return out


# ─── выгрузка ──────────────────────────────────────────────────────────────


def save_json(bid: Bid, path: Path) -> None:
    path.write_text(json.dumps(asdict(bid), ensure_ascii=False, indent=1), encoding="utf-8")


def save_csv(bid: Bid, path: Path) -> None:
    """Для переноса в систему: час по форме и объём в МВт с запятой."""
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(
            ["Операционные сутки", "Час", "Интервал (Астана)", "МВт", "P10, МВт", "P90, МВт"]
        )
        for h in bid.hours:
            w.writerow(
                [
                    bid.operational_day,
                    f"{h.hour:02d}:00",
                    h.interval_local,
                    _fmt(h.mw),
                    _fmt(h.p10_mw),
                    _fmt(h.p90_mw),
                ]
            )


def save_docx(bid: Bid, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    banner = doc.add_paragraph()
    run = banner.add_run(DRAFT_BANNER)
    run.bold = True
    run.font.color.rgb = RGBColor(0xB0, 0x1C, 0x1C)
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Приложение 7 к Правилам организации и функционирования оптового рынка "
        "электрической энергии"
    ).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title = doc.add_heading("Заявка на продажу", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    head = doc.add_table(rows=0, cols=2)
    head.style = "Table Grid"
    for k, v in (
        ("Отправитель", bid.sender),
        ("Контрагент", bid.counterparty),
        ("Операция", bid.operation),
        ("Операционные сутки", f"{pd.Timestamp(bid.operational_day):%d.%m.%Y}"),
        ("Срок подачи (п. 51)", f"до {bid.deadline_local} по времени Астаны"),
    ):
        row = head.add_row().cells
        row[0].text, row[1].text = k, v
        row[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    grid = doc.add_table(rows=13, cols=4)
    grid.style = "Table Grid"
    for c, text in enumerate(("Час", "Объём, МВт", "Час", "Объём, МВт")):
        grid.cell(0, c).text = text
        grid.cell(0, c).paragraphs[0].runs[0].bold = True
    for i in range(12):
        for j, h in enumerate((bid.hours[i], bid.hours[i + 12])):
            grid.cell(i + 1, 2 * j).text = f"{h.hour:02d}:00"
            grid.cell(i + 1, 2 * j + 1).text = _fmt(h.mw)
            grid.cell(i + 1, 2 * j + 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("Примечание: значения указываются с точностью до тысячных.")
    doc.add_paragraph(f"Итого за сутки: {_fmt(bid.total_mwh)} МВт·ч.").runs[0].bold = True

    doc.add_heading("Обоснование объёмов (справочно, в заявку не входит)", level=2)
    doc.add_paragraph(
        f"Прогноз {bid.forecast_id}, сформирован {bid.prepared_at_local} (Астана); модель "
        f"{bid.model_version}; погода: {bid.weather_provider}, выпуск {bid.weather_run} UTC."
    )
    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    for c, text in enumerate(("Час", "Интервал", "P10, МВт", "P50, МВт", "P90, МВт", "Ветер, м/с")):
        t.rows[0].cells[c].text = text
    for h in bid.hours:
        r = t.add_row().cells
        vals = (
            f"{h.hour:02d}:00",
            h.interval_local,
            _fmt(h.p10_mw),
            _fmt(h.mw),
            _fmt(h.p90_mw),
            f"{h.wind_ms:.1f}" + (" ❄" if h.icing else ""),
        )
        for c, v in enumerate(vals):
            r[c].text = v
    if bid.risks:
        doc.add_paragraph("Риски по оценке агента:").runs[0].bold = True
        for r in bid.risks:
            doc.add_paragraph(r, style="List Bullet")
    if bid.agent_summary:
        doc.add_paragraph("Сводка агента: " + bid.agent_summary)

    doc.add_heading("Черновики корректировок (п. 97–99)", level=2)
    if bid.corrections:
        doc.add_paragraph(
            "Подаются не позднее чем за 2 часа до часа; корректировка ВИЭ исполняется только "
            "при встречной корректировке другой ВИЭ, частичное исполнение допускается."
        )
        ct = doc.add_table(rows=1, cols=6)
        ct.style = "Table Grid"
        for c, text in enumerate(
            ("Решение агента", "Час", "Направление", "Объём, МВт", "Было → стало, МВт", "Подать до")
        ):
            ct.rows[0].cells[c].text = text
        for c in bid.corrections:
            r = ct.add_row().cells
            vals = (
                c.decided_at_local,
                f"{c.hour:02d}:00",
                c.direction,
                _fmt(c.volume_mw),
                f"{_fmt(c.was_mw)} → {_fmt(c.new_mw)}",
                c.submit_before_local,
            )
            for i, v in enumerate(vals):
                r[i].text = v
    else:
        doc.add_paragraph(
            "Ревизии прогноза не превысили порог корректировки — корректировки не нужны."
        )

    doc.add_heading("Допущения", level=2)
    for a in bid.assumptions:
        doc.add_paragraph(a, style="List Bullet")
    doc.save(path)


def _register_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONTS / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONTS / "DejaVuSans-Bold.ttf")))
    return "DejaVu", "DejaVu-Bold"


def save_pdf(bid: Bid, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _register_fonts()
    base = ParagraphStyle("b", fontName=font, fontSize=9, leading=12)
    small = ParagraphStyle(
        "s", parent=base, fontSize=8, leading=10, textColor=colors.HexColor("#444444")
    )
    right = ParagraphStyle("r", parent=small, alignment=TA_RIGHT)
    h1 = ParagraphStyle(
        "h1", parent=base, fontName=bold, fontSize=15, leading=19, alignment=TA_CENTER
    )
    h2 = ParagraphStyle("h2", parent=base, fontName=bold, fontSize=11, leading=14, spaceBefore=8)
    warn = ParagraphStyle(
        "w",
        parent=base,
        fontName=bold,
        textColor=colors.HexColor("#B01C1C"),
        alignment=TA_CENTER,
        borderColor=colors.HexColor("#B01C1C"),
        borderWidth=1,
        borderPadding=4,
    )
    grid_style = [
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#888888")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    head_style = grid_style + [
        ("FONTNAME", (0, 0), (-1, 0), bold),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF1F4")),
    ]

    def page(canvas, doc_):
        canvas.saveState()
        canvas.setFont(bold, 60)
        canvas.setFillColor(colors.Color(0.7, 0.1, 0.1, alpha=0.08))
        canvas.translate(A4[0] / 2, A4[1] / 2)
        canvas.rotate(35)
        canvas.drawCentredString(0, 0, "ЧЕРНОВИК")
        canvas.restoreState()
        canvas.setFont(font, 7)
        canvas.drawString(15 * mm, 10 * mm, f"windcast · {bid.forecast_id} · стр. {doc_.page}")

    story = [
        Paragraph(DRAFT_BANNER, warn),
        Spacer(1, 6),
        Paragraph(
            "Приложение 7 к Правилам организации и функционирования оптового рынка "
            "электрической энергии",
            right,
        ),
        Spacer(1, 4),
        Paragraph("Заявка на продажу", h1),
        Spacer(1, 8),
    ]
    head = [
        ["Отправитель", Paragraph(bid.sender, base)],
        ["Контрагент", Paragraph(bid.counterparty, base)],
        ["Операция", bid.operation],
        ["Операционные сутки", f"{pd.Timestamp(bid.operational_day):%d.%m.%Y}"],
        ["Срок подачи (п. 51)", f"до {bid.deadline_local} по времени Астаны"],
    ]
    ht = Table(head, colWidths=[45 * mm, 125 * mm])
    ht.setStyle(TableStyle(grid_style + [("FONTNAME", (0, 0), (0, -1), bold)]))
    story += [ht, Spacer(1, 8)]

    rows = [["Час", "Объём, МВт", "Час", "Объём, МВт"]]
    for i in range(12):
        a, b = bid.hours[i], bid.hours[i + 12]
        rows.append([f"{a.hour:02d}:00", _fmt(a.mw), f"{b.hour:02d}:00", _fmt(b.mw)])
    gt = Table(rows, colWidths=[30 * mm, 55 * mm, 30 * mm, 55 * mm])
    gt.setStyle(
        TableStyle(
            head_style + [("ALIGN", (1, 1), (1, -1), "RIGHT"), ("ALIGN", (3, 1), (3, -1), "RIGHT")]
        )
    )
    story += [
        gt,
        Spacer(1, 4),
        Paragraph("Примечание: значения указываются с точностью до тысячных.", small),
        Paragraph(f"<font name='{bold}'>Итого за сутки: {_fmt(bid.total_mwh)} МВт·ч</font>", base),
        Spacer(1, 6),
    ]

    story.append(Paragraph("Обоснование объёмов (справочно, в заявку не входит)", h2))
    story.append(
        Paragraph(
            f"Прогноз {bid.forecast_id}, сформирован {bid.prepared_at_local} (Астана); модель "
            f"{bid.model_version}; погода: {bid.weather_provider}, выпуск {bid.weather_run} UTC.",
            small,
        )
    )
    rows = [["Час", "Интервал", "P10, МВт", "P50, МВт", "P90, МВт", "Ветер, м/с"]]
    for h in bid.hours:
        rows.append(
            [
                f"{h.hour:02d}:00",
                h.interval_local,
                _fmt(h.p10_mw),
                _fmt(h.mw),
                _fmt(h.p90_mw),
                f"{h.wind_ms:.1f}" + (" *" if h.icing else ""),
            ]
        )
    jt = Table(rows, colWidths=[18 * mm, 30 * mm, 28 * mm, 28 * mm, 28 * mm, 26 * mm], repeatRows=1)
    jt.setStyle(
        TableStyle(
            head_style + [("FONTSIZE", (0, 0), (-1, -1), 8), ("ALIGN", (2, 1), (-1, -1), "RIGHT")]
        )
    )
    story += [Spacer(1, 4), jt, Paragraph("* риск обледенения", small)]
    if bid.risks:
        story.append(Paragraph("Риски по оценке агента:", h2))
        story += [Paragraph("• " + r, base) for r in bid.risks]
    if bid.agent_summary:
        story += [Spacer(1, 4), Paragraph("Сводка агента: " + bid.agent_summary, small)]

    story.append(Paragraph("Черновики корректировок (п. 97–99)", h2))
    if bid.corrections:
        story.append(
            Paragraph(
                "Подаются не позднее чем за 2 часа до часа; корректировка ВИЭ исполняется "
                "только при встречной корректировке другой ВИЭ, частичное исполнение допускается.",
                small,
            )
        )
        rows = [
            ["Решение агента", "Час", "Направление", "Объём, МВт", "Было → стало, МВт", "Подать до"]
        ]
        for c in bid.corrections:
            rows.append(
                [
                    c.decided_at_local,
                    f"{c.hour:02d}:00",
                    c.direction,
                    _fmt(c.volume_mw),
                    f"{_fmt(c.was_mw)} → {_fmt(c.new_mw)}",
                    c.submit_before_local,
                ]
            )
        ct = Table(
            rows, colWidths=[30 * mm, 14 * mm, 24 * mm, 22 * mm, 42 * mm, 28 * mm], repeatRows=1
        )
        ct.setStyle(TableStyle(head_style + [("FONTSIZE", (0, 0), (-1, -1), 8)]))
        story += [Spacer(1, 4), ct]
    else:
        story.append(
            Paragraph("Ревизии прогноза не превысили порог — корректировки не нужны.", base)
        )

    story.append(Paragraph("Допущения", h2))
    story += [Paragraph("• " + a, small) for a in bid.assumptions]

    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=15 * mm,
        title=f"Заявка на продажу {bid.operational_day} (черновик)",
    ).build(story, onFirstPage=page, onLaterPages=page)


def export(
    operational_day: str,
    sender: str = DEFAULT_SENDER,
    rated_mw_per_turbine: float = RATED_MW_PER_TURBINE,
) -> dict:
    bid = build_bid(operational_day, sender, rated_mw_per_turbine)
    out = BIDS_DIR / bid.operational_day
    out.mkdir(parents=True, exist_ok=True)
    stem = f"zayavka_{bid.operational_day}"
    files = {
        "json": out / f"{stem}.json",
        "csv": out / f"{stem}.csv",
        "docx": out / f"{stem}.docx",
        "pdf": out / f"{stem}.pdf",
    }
    save_json(bid, files["json"])
    save_csv(bid, files["csv"])
    save_docx(bid, files["docx"])
    save_pdf(bid, files["pdf"])
    return {
        "day": bid.operational_day,
        "total_mwh": bid.total_mwh,
        "corrections": len(bid.corrections),
        **{k: str(v) for k, v in files.items()},
    }


def export_all(first: str = "2026-02-01", last: str = "2026-02-28", **kw) -> list[dict]:
    res = [export(str(d.date()), **kw) for d in pd.date_range(first, last, freq="D")]
    (BIDS_DIR / "index.json").write_text(
        json.dumps(
            [{k: r[k] for k in ("day", "total_mwh", "corrections")} for r in res],
            ensure_ascii=False,
            indent=1,
        ),
        encoding="utf-8",
    )
    return res
